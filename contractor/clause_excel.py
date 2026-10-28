"""Shared clause spreadsheet format and validation for CR and Tender Analysis."""
import io
import zipfile

import frappe

from contractor.tender_logic import number_rows

HEADERS = ["Series Number", "Contractual Clause Title", "Is Group", "Qty", "UOM"]


def read_clause_excel(file_url):
    from openpyxl import load_workbook
    file = frappe.get_doc("File", {"file_url": file_url})
    file.check_permission("read")
    content = file.get_content()
    if len(content) > 5 * 1024 * 1024:
        frappe.throw("Excel file must be no larger than 5 MB")
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            if sum(entry.file_size for entry in archive.infolist()) > 30 * 1024 * 1024:
                frappe.throw("Excel workbook is too large when expanded")
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=False)
        sheet = workbook.active
        records = sheet.iter_rows(values_only=True)
        if list(next(records, ()))[:5] != HEADERS:
            frappe.throw("Use the provided Excel template and its column headings")
        rows = []
        for index, values in enumerate(records, 2):
            if index > 10002:
                frappe.throw("Maximum 10,000 Excel rows")
            if not any(v is not None for v in values):
                continue
            values = list(values[:5]) + [None] * max(0, 5-len(values))
            series, title, flag, qty, uom = values
            if any(isinstance(v, str) and v.startswith("=") for v in values):
                frappe.throw(f"Excel row {index}: formulas are not supported")
            if str(flag or 0).lower() not in ("0", "1", "yes", "no", "true", "false"):
                frappe.throw(f"Excel row {index}: Is Group must be 0 or 1")
            if isinstance(series, (float, int)) and float(series).is_integer():
                series = str(int(series))
            row = dict(series_number=str(series or ""), clause_title=str(title or ""),
                is_group=int(str(flag).lower() in ("1", "yes", "true")), qty=float(qty or 0), uom=uom)
            if uom and not frappe.db.exists("UOM", uom):
                frappe.throw(f"Excel row {index}: unknown UOM {uom}")
            rows.append(row)
        number_rows(rows, "From Excel")
        if not rows:
            frappe.throw("The spreadsheet has no clauses")
        return rows
    except (ValueError, TypeError, zipfile.BadZipFile) as exc:
        frappe.throw(str(exc))
    finally:
        if "workbook" in locals():
            workbook.close()

