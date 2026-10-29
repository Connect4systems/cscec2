"""Build downloadable Arabic/English DOCX, printable HTML and PDF guides."""
from pathlib import Path
from html import escape
import json
import re
import zipfile
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from playwright.sync_api import sync_playwright
from guide_content import SECTIONS

ROOT = Path(__file__).resolve().parent
OUT = ROOT / 'downloads'
OUT.mkdir(exist_ok=True)
DATE = '2026-10-29'
AR = re.compile(r'[\u0600-\u06ff]')


def words(par, text, rtl):
    par.paragraph_format.space_after = Pt(7)
    par.paragraph_format.line_spacing = 1.15
    if rtl:
        prop = par._p.get_or_add_pPr()
        bidi = OxmlElement('w:bidi'); prop.append(bidi)
        par.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for part in re.split(r'([A-Za-z0-9][A-Za-z0-9_:/?%#&=<>.,+×÷−()’\- ]*)', text):
        if not part:
            continue
        run = par.add_run(part)
        run.font.name = 'Arial'
        run.font.size = Pt(11)
        rpr = run._element.get_or_add_rPr()
        fonts = rpr.rFonts
        if fonts is not None:
            fonts.set(qn('w:cs'), 'Arial')
        if AR.search(part):
            direction = OxmlElement('w:rtl'); rpr.append(direction)
            size = OxmlElement('w:szCs'); size.set(qn('w:val'), '24'); rpr.append(size)
    return par


def caption(filename, lang):
    label = filename[:-4].replace('-', ' ').title()
    return ('صورة فعلية من النظام: ' if lang == 'ar' else 'System screenshot: ') + label


def site_note(lang):
    if lang == 'ar':
        return 'أُعد هذا الدليل من ملفات التطبيق مع مراجعة شاشات الموقع بحساب مصرح به. الصور من النظام الفعلي، وقد تعرض بيانات تجريبية موجودة. شملت المراجعة تسجيل الدخول وفتح النماذج والتبويبات وتجربة النوافذ والحسابات. لم تُنشأ أو تُعتمد فواتير أو مستخلصات مالية. استُخدمت مسودة CR مؤقتة لاختبار الجلب وحُذفت بعده، وقد تقدم تسلسل التسمية. تختلف الأزرار والقيم بحسب صلاحياتك وإعدادات الموقع. لا يتضمن الدليل بيانات الدخول.'
    return 'This guide combines application source review with authorized inspection of the live site. Screenshots are from the actual system and may show existing test data. Review covered sign-in, forms, tabs and selected dialogs/browser calculations. No invoices or financial certificates were created or submitted for this guide. A temporary CR draft was used to test fetching and deleted afterwards; its naming sequence may advance. Buttons and values can differ with your permissions and site settings. Login credentials are excluded.'


def validation_text(lang):
    path = ROOT / 'site-review.json'
    d = json.loads(path.read_text(encoding='utf-8')) if path.exists() else {'results': []}
    results = d['results']
    ok = sum(x['status'] == 'opened' for x in results)
    errors = [x['doctype'] + ': ' + x.get('error', '')[:200] for x in results if x['status'] != 'opened']
    text = (f'تم فتح {ok} من {len(results)} نوع مستند خلال مراجعة الموقع. نجحت الاختبارات المحلية البالغ عددها 172. لا تمثل هذه النتائج اختبار ترحيل محاسبي كامل من البداية للنهاية. ' if lang == 'ar' else f'{ok} of {len(results)} document types opened during site review. All 172 local automated tests passed. These checks do not establish end-to-end accounting posting success. ')
    if errors:
        text += ('عناصر تحتاج مراجعة: ' if lang == 'ar' else 'Items requiring review: ') + '; '.join(errors)
    browser_errors = d.get('browser_errors', [])
    if browser_errors:
        text += (' لوحظ تنبيه JavaScript أثناء التنقل يحتاج إلى تتبع مستقل: ' if lang == 'ar' else ' A JavaScript message observed during navigation needs separate investigation: ') + '; '.join(str(x) for x in browser_errors)
    return text


def build(lang):
    rtl = lang == 'ar'
    title = 'دليل مستخدم تطبيق المقاولات' if rtl else 'Contractor App User Guide'
    subtitle = 'دليل عملي خطوة بخطوة — دورة المشروع والموازنة ومقاولي الباطن' if rtl else 'Step-by-step project, budget and subcontracting workflows'
    doc = Document()
    section = doc.sections[0]
    section.top_margin = section.bottom_margin = Inches(.65)
    section.left_margin = section.right_margin = Inches(.65)
    section.page_width = Inches(8.27); section.page_height = Inches(11.69)
    for style in ['Normal', 'Heading 1', 'Heading 2', 'Title', 'Subtitle']:
        doc.styles[style].font.name = 'Arial'
    doc.styles['Normal'].font.size = Pt(11)
    doc.styles['Heading 1'].font.color.rgb = RGBColor.from_string('123F54')
    header = section.header.paragraphs[0]
    words(header, 'Contractor App | ' + ('دليل المستخدم' if rtl else 'User Guide'), rtl)
    header.runs[0].font.size = Pt(9)
    foot = section.footer.paragraphs[0]
    words(foot, DATE + '  |  ', rtl)
    field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE'); foot._p.append(field)
    words(doc.add_paragraph(style='Title'), title, rtl)
    words(doc.add_paragraph(style='Subtitle'), subtitle, rtl)
    words(doc.add_paragraph(), 'en.connect4systems.com | ' + DATE, rtl)
    words(doc.add_paragraph(), site_note(lang), rtl)
    words(doc.add_paragraph(), 'ERPNext 15 · Frappe 15 · Contractor App', rtl)
    doc.add_page_break()
    contents = 'المحتويات' if rtl else 'Contents'
    words(doc.add_paragraph(style='Heading 1'), contents, rtl)
    for i, s in enumerate(SECTIONS, 1):
        words(doc.add_paragraph(), f'{i:02d}. ' + s[lang], rtl)
    html = [f'<!doctype html><html lang="{lang}" dir="{"rtl" if rtl else "ltr"}"><head><meta charset="utf-8"><title>{title}</title><style>' + CSS + '</style></head><body>',
            f'<div class="cover"><div class="brand">CONTRACTOR APP</div><h1>{title}</h1><h2>{subtitle}</h2><p>en.connect4systems.com · {DATE}</p><aside>{escape(site_note(lang))}</aside><p>ERPNext 15 · Frappe 15 · Contractor App</p></div>',
            f'<section class="contents"><h1>{contents}</h1><ol>']
    for i, s in enumerate(SECTIONS, 1):
        html.append(f'<li><a href="#chapter-{i}">{escape(s[lang])}</a></li>')
    html.append('</ol></section>')
    for i, s in enumerate(SECTIONS, 1):
        doc.add_page_break()
        words(doc.add_paragraph(style='Heading 1'), f'{i:02d}. ' + s[lang], rtl)
        html.append(f'<section id="chapter-{i}"><div class="eyebrow">{i:02d} / {len(SECTIONS):02d}</div><h1>{escape(s[lang])}</h1><ol class="steps">')
        for j, step in enumerate(s['steps_' + lang], 1):
            words(doc.add_paragraph(), f'{i}.{j}  ' + step, rtl)
            html.append(f'<li><span class="step-no">{i}.{j}</span><div>{escape(step)}</div></li>')
        html.append('</ol>')
        if s['note_' + lang]:
            words(doc.add_paragraph(), ('مراجعة مهمة: ' if rtl else 'Check: ') + s['note_' + lang], rtl)
            html.append('<aside>' + escape(s['note_' + lang]) + '</aside>')
        images = [name for name in s['images'] if (ROOT / 'assets' / name).exists()]
        for name in images:
            doc.add_page_break()
            words(doc.add_paragraph(style='Heading 2'), s[lang], rtl)
            doc.add_picture(str(ROOT / 'assets' / name), width=Inches(6.9))
            words(doc.add_paragraph(), caption(name, lang), rtl)
            html.append(f'<figure><img src="../assets/{name}" alt="{escape(caption(name,lang))}"><figcaption>{escape(caption(name,lang))}</figcaption></figure>')
        html.append('</section>')
    doc.add_page_break()
    words(doc.add_paragraph(style='Heading 1'), 'نطاق التحقق من الموقع' if rtl else 'Site verification scope', rtl)
    words(doc.add_paragraph(), validation_text(lang), rtl)
    html.append('<section><h1>' + ('نطاق التحقق من الموقع' if rtl else 'Site verification scope') + '</h1><p>' + escape(validation_text(lang)) + '</p><p>' + escape(site_note(lang)) + '</p>')
    interaction = ROOT / 'interaction-review.json'
    if interaction.exists():
        checks = json.loads(interaction.read_text(encoding='utf-8'))['checks']
        for item in checks:
            state = ('ناجح' if rtl else 'Passed') if item.get('passed') else item.get('status', item.get('error', 'Review'))
            line = item['check'] + ' — ' + state
            words(doc.add_paragraph(), line, rtl)
            html.append('<p dir="ltr">' + escape(line) + '</p>')
    html.append('</section>')
    html.append('</body></html>')
    base = OUT / ('Contractor-User-Guide-' + lang.upper())
    doc.save(base.with_suffix('.docx'))
    base.with_suffix('.html').write_text('\n'.join(html), encoding='utf-8')
    return base


CSS = """
@page { size: A4; margin: 16mm 15mm 18mm; }
* { box-sizing:border-box; }
body { margin:0; color:#21343e; font:11pt/1.7 'Segoe UI',Arial,sans-serif; }
h1,h2 { color:#123f54; line-height:1.35; }
h1 { font-size:23pt; margin:4mm 0 7mm; }
h2 { font-size:16pt; }
a { color:#12647b; text-decoration:none; }
.cover { min-height:245mm; padding-top:35mm; break-after:page; }
.cover h1 {font-size:36pt; max-width:160mm;}
.brand,.eyebrow {letter-spacing:1px; color:#008684; font-size:11pt; font-weight:700;}
.contents {break-after:page;}
.contents li {margin:1mm 0;}
section {break-before:page;}
.steps {list-style:none;padding:0;}
.steps li {display:flex; gap:4mm; margin-bottom:4mm; break-inside:avoid;}
.step-no {min-width:11mm; color:#008684; font-weight:700;direction:ltr;}
aside {background:#edf5f6; border-inline-start:3px solid #008684; padding:4mm; margin:5mm 0; break-inside:avoid;}
figure {margin:0; break-before:page; break-inside:avoid; padding-top:8mm;}
img {display:block; width:100%; max-height:233mm; object-fit:contain; object-position:top; border:1px solid #d5e0e4;}
figcaption {font-size:9pt; margin-top:4mm;color:#526c79;}
@media screen {body {max-width:900px;margin:auto;padding:30px;} section {margin-top:70px;} figure {margin-top:30px;}}
"""


def main():
    bases = [build('ar'), build('en')]
    with sync_playwright() as p:
        browser = p.chromium.launch(executable_path=r'C:/Program Files (x86)/Microsoft/Edge/Application/msedge.exe', headless=True)
        page = browser.new_page()
        for base in bases:
            page.goto(base.with_suffix('.html').as_uri(), wait_until='networkidle')
            page.pdf(path=str(base.with_suffix('.pdf')), print_background=True, prefer_css_page_size=True,
                     display_header_footer=True, header_template='<span></span>',
                     footer_template='<div style="font-size:9px;width:100%;text-align:center;color:#667;">Contractor App · 2026-10-29 · <span class="pageNumber"></span> / <span class="totalPages"></span></div>')
            print(base.name, 'DOCX / HTML / PDF built', flush=True)
        browser.close()
    archive = OUT / 'Contractor-User-Guides-AR-EN.zip'
    with zipfile.ZipFile(archive, 'w', zipfile.ZIP_DEFLATED) as z:
        for base in bases:
            for ext in ['.pdf', '.docx', '.html']:
                z.write(base.with_suffix(ext), 'downloads/' + base.with_suffix(ext).name)
        for img in sorted((ROOT/'assets').glob('*.png')):
            z.write(img, 'assets/' + img.name)
        for name in ['README.md', 'site-verification.md']:
            if (ROOT/name).exists(): z.write(ROOT/name,name)
    print('ZIP:', archive, flush=True)


if __name__ == '__main__':
    main()
